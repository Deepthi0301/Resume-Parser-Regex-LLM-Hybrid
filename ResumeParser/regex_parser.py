
import re
import json
import sys
from pathlib import Path

import fitz
from docx import Document
from odf.opendocument import load
from odf import text, teletype


def extract_text(file_path):
    """Read text from PDF, DOCX or ODT files."""
    path = Path(file_path)
    extension = path.suffix.lower()

    if not path.is_file():
        raise FileNotFoundError("Resume file not found.")

    if extension == ".pdf":
        with fitz.open(str(path)) as pdf:
            return "\n".join(page.get_text() for page in pdf)

    elif extension == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)

        return "\n".join(parts)

    elif extension == ".odt":
        doc = load(str(path))
        paragraphs = doc.getElementsByType(text.P)
        return "\n".join(teletype.extractText(p) for p in paragraphs)

    else:
        raise ValueError(
            "Unsupported file format. Use PDF, DOCX or ODT."
        )


def extract_information(resume_text):
    """Extract candidate details using regex and simple rules."""

    # Email address
    email_match = re.search(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        resume_text
    )
    email = email_match.group(0) if email_match else None

    # Phone number
    phone_match = re.search(
        r"(?<!\w)(?:\+?\d[\d\s().-]{7,}\d)(?!\w)",
        resume_text
    )
    phone = phone_match.group(0).strip() if phone_match else None

    # Full name: check the first few non-empty lines
    lines = [line.strip() for line in resume_text.splitlines() if line.strip()]
    name = None

    for line in lines[:5]:
        if (
            len(line.split()) <= 5
            and re.fullmatch(r"[A-Za-z]+(?:[ .'-][A-Za-z]+)*", line)
            and not re.search(
                r"resume|curriculum vitae|email|phone|profile",
                line,
                re.IGNORECASE
            )
        ):
            name = line
            break

    # Technical skills
    known_skills = [
        "Python", "Java", "JavaScript", "SQL", "HTML", "CSS",
        "React", "Node.js", "Machine Learning", "Data Analytics",
        "ServiceNow", "AWS", "Azure", "Git", "Excel", "Power BI",
        "C++", "C", "Linux", "REST API"
    ]

    skills = []
    for skill in known_skills:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
        if re.search(pattern, resume_text, re.IGNORECASE):
            skills.append(skill)

    # Recognized resume section headings
    all_headings = {
        "EDUCATION",
        "EDUCATIONAL QUALIFICATIONS",
        "ACADEMIC QUALIFICATIONS",
        "WORK EXPERIENCE",
        "PROFESSIONAL EXPERIENCE",
        "EXPERIENCE",
        "INTERNSHIP EXPERIENCE",
        "TECHNICAL SKILLS",
        "SKILLS",
        "PROJECTS",
        "ACADEMIC PROJECTS",
        "CERTIFICATIONS",
        "CERTIFICATES",
        "PERSONAL DETAILS",
        "PERSONAL INFORMATION",
        "ACHIEVEMENTS",
        "LANGUAGES",
        "INTERESTS",
        "HOBBIES",
        "SUMMARY",
        "PROFESSIONAL SUMMARY",
        "CAREER OBJECTIVE",
        "OBJECTIVE",
        "DECLARATION",
        "PUBLICATIONS",
        "POSITIONS OF RESPONSIBILITY"
    }

    def normalize_heading(line):
        """Normalize a heading for comparison."""
        return re.sub(
            r"[^A-Z0-9 ]", "", line.upper()
        ).strip()

    def get_section(headings):
        """Extract section text until the next recognized heading."""
        target_headings = {
            normalize_heading(heading) for heading in headings
        }

        collecting = False
        content = []

        for line in resume_text.splitlines():
            cleaned = line.strip()

            if not cleaned:
                if collecting:
                    content.append("")
                continue

            normalized = normalize_heading(cleaned)

            # Start collecting after the requested heading
            if normalized in target_headings:
                collecting = True
                continue

            # Stop when another section heading is reached
            if collecting and normalized in all_headings:
                break

            # Collect the section's content
            if collecting:
                content.append(cleaned)

        result = "\n".join(content).strip()
        return result if result else None

    # Extract sections
    education = get_section([
        "EDUCATION",
        "EDUCATIONAL QUALIFICATIONS",
        "ACADEMIC QUALIFICATIONS"
    ])

    experience = get_section([
        "WORK EXPERIENCE",
        "PROFESSIONAL EXPERIENCE",
        "EXPERIENCE",
        "INTERNSHIP EXPERIENCE"
    ])

    projects = get_section([
        "PROJECTS",
        "ACADEMIC PROJECTS"
    ])

    certifications = get_section([
        "CERTIFICATIONS",
        "CERTIFICATES"
    ])

    return {
        "full_name": name,
        "email_address": email,
        "phone_number": phone,
        "technical_skills": skills,
        "educational_qualifications": education,
        "work_experience": experience,
        "projects": projects if projects else [],
        "certifications": certifications if certifications else []
    }


def parse_resume(file_path):
    """Parse a resume and save the extracted information as JSON."""
    try:
        resume_text = extract_text(file_path)

        if not resume_text.strip():
            print("Error: No text found in this resume.")
            return

        result = extract_information(resume_text)

        output_dir = Path("output")
        output_dir.mkdir(exist_ok=True)

        output_file = output_dir / (
            Path(file_path).stem + "_parsed.json"
        )

        output_file.write_text(
            json.dumps(result, indent=4, ensure_ascii=False),
            encoding="utf-8"
        )

        print(json.dumps(result, indent=4, ensure_ascii=False))
        print(f"\nJSON saved to: {output_file.resolve()}")

    except (FileNotFoundError, ValueError) as error:
        print(f"Error: {error}")
    except Exception as error:
        print(f"Could not parse resume: {error}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python regex_parser.py <resume_file>")
    else:
        parse_resume(sys.argv[1])
