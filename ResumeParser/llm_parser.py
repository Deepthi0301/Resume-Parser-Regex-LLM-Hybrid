
import json
import sys
import urllib.request
import urllib.error
from pathlib import Path

import fitz
from docx import Document
from odf.opendocument import load
from odf import text, teletype

from regex_parser import extract_information as regex_extract


# Step 1: Read text from PDF, DOCX or ODT
def extract_text(file_path):
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError("Resume file not found.")

    extension = path.suffix.lower()

    if extension == ".pdf":
        with fitz.open(str(path)) as pdf:
            return "\n".join(page.get_text() for page in pdf)

    elif extension == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)

        return "\n".join(parts)

    elif extension == ".odt":
        doc = load(str(path))
        paragraphs = doc.getElementsByType(text.P)
        return "\n".join(
            teletype.extractText(p) for p in paragraphs
        )

    else:
        raise ValueError(
            "Unsupported format. Use PDF, DOCX or ODT."
        )


# Step 2: Extract information using the local AI model
def extract_information(resume_text):
    fields = {
        "full_name": None,
        "email_address": None,
        "phone_number": None,
        "technical_skills": [],
        "educational_qualifications": None,
        "work_experience": None,
        "projects": [],
        "certifications": []
    }

    prompt = """
You are a resume information extraction assistant.

Read the COMPLETE resume and extract all available information.
Never invent details. Do not skip sections containing information.

Return valid JSON only with exactly these keys:
full_name, email_address, phone_number, technical_skills,
educational_qualifications, work_experience, projects,
certifications.

Rules:
1. Text fields must be strings or null.
2. technical_skills, projects and certifications must be lists of strings.
3. Extract every education entry, including degree, institution,
   year and CGPA or grade when available.
4. Extract every job and internship, including role, company,
   dates and responsibilities when available.
5. Include every project and its description.
6. Return individual skill names, not grouped categories.
7. Use null for missing text fields and [] for missing lists.

RESUME TEXT:
""" + resume_text + """

Return JSON only.
"""

    payload = {
        "model": "llama3.2:1b",
        "prompt": prompt,
        "format": "json",
        "stream": False,
        "options": {
            "temperature": 0,
            "num_predict": 1500
        }
    }

    request = urllib.request.Request(
        "http://localhost:11434/api/generate",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST"
    )

    with urllib.request.urlopen(request, timeout=300) as response:
        result = json.loads(response.read().decode("utf-8"))

    print("\nAI raw response:")
    print(result.get("response", ""))

    extracted = json.loads(result["response"])

    # Keep the required fields and validate their types
    final_result = {}

    for key, default in fields.items():
        value = extracted.get(key, default)

        if key in (
            "technical_skills",
            "projects",
            "certifications"
        ):
            if not isinstance(value, list):
                value = []

            value = [
                str(item).strip()
                for item in value
                if isinstance(item, (str, int, float))
                and str(item).strip()
            ]

        elif key in (
            "full_name",
            "email_address",
            "phone_number",
            "educational_qualifications",
            "work_experience"
        ):
            if not isinstance(value, str) or not value.strip():
                value = None

        final_result[key] = value

    return final_result


# Step 3: Parse resume and use regex as a fallback
def parse_resume(file_path):
    try:
        resume_text = extract_text(file_path)

        if not resume_text.strip():
            print("Error: No text found in this resume.")
            return

        # First, extract using AI
        result = extract_information(resume_text)

        # Then, extract using regex as a fallback
        regex_data = regex_extract(resume_text)

        # Fill missing text fields with regex results
        text_fields = [
            "full_name",
            "email_address",
            "phone_number",
            "educational_qualifications",
            "work_experience"
        ]

        for field in text_fields:
            if not result.get(field) and regex_data.get(field):
                result[field] = regex_data[field]

        # Use regex lists if the AI returned empty lists
        list_fields = [
            "technical_skills",
            "projects",
            "certifications"
        ]

        for field in list_fields:
            if not result.get(field) and regex_data.get(field):
                value = regex_data[field]

                if isinstance(value, list):
                    result[field] = value
                elif isinstance(value, str):
                    result[field] = [value]

        # Save the JSON output
        output_dir = Path("output")
        output_dir.mkdir(exist_ok=True)

        output_file = output_dir / (
            Path(file_path).stem + "_llm_parsed.json"
        )

        output_file.write_text(
            json.dumps(result, indent=4, ensure_ascii=False),
            encoding="utf-8"
        )

        print("\nFinal extracted information:")
        print(json.dumps(result, indent=4, ensure_ascii=False))
        print(f"\nJSON saved to: {output_file.resolve()}")

    except FileNotFoundError as error:
        print(f"Error: {error}")

    except ValueError as error:
        print(f"Error: {error}")

    except urllib.error.URLError:
        print(
            "Error: Cannot connect to Ollama. "
            "Make sure Ollama is running locally."
        )

    except json.JSONDecodeError:
        print("Error: The AI returned invalid JSON.")

    except Exception as error:
        print(f"Could not parse resume: {error}")


# Step 4: Run the program
if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python llm_parser.py <resume_file>")
    else:
        parse_resume(sys.argv[1])
